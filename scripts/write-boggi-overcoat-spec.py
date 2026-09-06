#!/usr/bin/env python3
"""Write Boggi Measurement Spec for Overcoat.xlsx (Format A)."""
import json
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip3 install openpyxl")

payload = json.load(sys.stdin)
sizes = payload["sizes"]
points = payload["points"]
out_path = payload["out_path"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Overcoat"
ws["A1"] = "Boggi Overcoat"
ws["A2"] = "Name"
ws["B2"] = "Boggi Overcoat"
ws["A3"] = "Code"
ws["B3"] = "BO-OC"
ws["A5"] = "Measurement point"
for index, size in enumerate(sizes, start=2):
    ws.cell(5, index, size)
ws.cell(5, len(sizes) + 2, "Remarks")

for row_index, point in enumerate(points, start=6):
    ws.cell(row_index, 1, point["name"])
    values = point.get("values") or {}
    for col_index, size in enumerate(sizes, start=2):
        value = values.get(size)
        if value is not None:
            ws.cell(row_index, col_index, value)
    if point.get("remark"):
        ws.cell(row_index, len(sizes) + 2, point["remark"])

wb.save(out_path)
print(out_path)
